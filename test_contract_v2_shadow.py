from __future__ import annotations

import ast
import copy
import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from shared_core.contract_v2_shadow import (
    build_question_contract_v2_shadow,
    write_question_contract_v2_shadow,
)
from shared_core.models import DocNode, QuestionUnit


PROJECT_ROOT = Path(__file__).resolve().parent
DOCUMENT_SHA256 = hashlib.sha256("影子原题".encode("utf-8")).hexdigest()


def _unit(question_id: str, source_span: tuple[int, int]) -> QuestionUnit:
    return QuestionUnit(
        question_id=question_id,
        subject="文科",
        subject_overlay="chinese",
        grade_hint="高三",
        question_type="choice",
        stem_blocks=[f"{question_id}．题干"],
        option_blocks=["A．甲", "B．乙"],
        subquestions=[],
        media_blocks=[],
        material_blocks=[],
        source_span=source_span,
        confidence=0.9,
        warnings=[],
    )


def test_shadow_adapter_keeps_legacy_objects_and_distinguishes_duplicate_labels() -> None:
    nodes = [
        DocNode(1, "1．第一篇第一题", metadata={"source_paragraph_index": 1}),
        DocNode(2, "A．甲", metadata={"source_paragraph_index": 2}),
        DocNode(3, "1．第二篇第一题", metadata={"source_paragraph_index": 3}),
        DocNode(4, "A．乙", metadata={"source_paragraph_index": 4}),
    ]
    units = [_unit("1", (1, 2)), _unit("1", (3, 4))]
    original_units = copy.deepcopy(units)
    original_nodes = copy.deepcopy(nodes)

    bundle = build_question_contract_v2_shadow(DOCUMENT_SHA256, units, nodes)

    assert units == original_units
    assert nodes == original_nodes
    contract = bundle["contract"]
    diff = bundle["diff"]
    assert contract["mode"] == "shadow"
    assert contract["production_execution_enabled"] is False
    assert contract["keypress_count"] == 0
    assert contract["blocking_issues"][0]["code"] == "SHADOW_CONTRACT_NOT_PRODUCTION"
    stable_ids = [unit["stable_unit_id"]["value"] for unit in contract["units"]]
    assert len(set(stable_ids)) == 2
    assert diff["duplicate_original_labels"][0]["label"] == "1"
    assert diff["duplicate_original_labels"][0]["count"] == 2
    assert diff["legacy_unit_count"] == diff["v2_unit_count"] == 2
    assert any("metadata" in item for item in diff["limitations"])


def test_virtual_segments_from_one_paragraph_receive_stable_occurrences() -> None:
    nodes = [
        DocNode(1, "1．第一问", metadata={"source_paragraph_index": 7}),
        DocNode(2, "2．同段拆出的第二问", metadata={"source_paragraph_index": 7}),
    ]
    units = [_unit("1", (1, 1)), _unit("2", (2, 2))]

    bundle = build_question_contract_v2_shadow(DOCUMENT_SHA256, units, nodes)
    refs = [
        unit["source_partition"]["included_refs"][0]
        for unit in bundle["contract"]["units"]
    ]

    assert [ref["native_id"] for ref in refs] == ["paragraph:7", "paragraph:7"]
    assert [ref["occurrence"] for ref in refs] == [1, 2]
    assert refs[0] != refs[1]


def test_media_metadata_remains_a_distinct_native_reference() -> None:
    media_sha = hashlib.sha256(b"image").hexdigest()
    nodes = [
        DocNode(
            1,
            "1．图示题",
            has_inline_media=True,
            metadata={"source_paragraph_index": 1, "media_sha256": [media_sha]},
        )
    ]

    bundle = build_question_contract_v2_shadow(
        DOCUMENT_SHA256,
        [_unit("1", (1, 1))],
        nodes,
    )
    refs = bundle["contract"]["units"][0]["source_partition"]["included_refs"]

    assert [ref["native_kind"] for ref in refs] == ["paragraph", "media"]
    assert bundle["diff"]["comparisons"][0]["native_kind_counts"] == {
        "formula": 0,
        "media": 1,
        "paragraph": 1,
        "table": 0,
    }


def test_shadow_contract_is_byte_stable_for_same_input() -> None:
    nodes = [DocNode(1, "1．题干", metadata={"source_paragraph_index": 1})]
    units = [_unit("1", (1, 1))]

    first = build_question_contract_v2_shadow(DOCUMENT_SHA256, units, nodes)
    second = build_question_contract_v2_shadow(DOCUMENT_SHA256, units, nodes)

    assert json.dumps(first, ensure_ascii=False, sort_keys=True) == json.dumps(
        second,
        ensure_ascii=False,
        sort_keys=True,
    )


def test_shadow_modules_do_not_import_execution_dependencies() -> None:
    forbidden = {
        "pyautogui",
        "win32com",
        "wps_helper",
        "answer_input",
        "core_parser",
    }
    for relative_path in (
        "shared_core/contracts_v2.py",
        "shared_core/contract_v2_shadow.py",
        "tools/build_question_contract_v2_shadow.py",
    ):
        path = PROJECT_ROOT / relative_path
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        imported = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.update(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.add(node.module)
        assert not {
            name
            for name in imported
            if any(part in forbidden for part in name.split("."))
        }, relative_path


def test_write_shadow_artifacts_is_readonly_and_outside_source_directory(
    tmp_path: Path,
) -> None:
    source_dir = tmp_path / "source"
    output_dir = tmp_path / "artifacts"
    source_dir.mkdir()
    source = source_dir / "重复题号.docx"
    document = Document()
    document.add_paragraph("1．第一篇第一题")
    document.add_paragraph("A．甲")
    document.add_paragraph("1．第二篇第一题")
    document.add_paragraph("A．乙")
    document.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()

    paths = write_question_contract_v2_shadow(source, output_dir)

    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert set(paths) == {"contract_json", "diff_json"}
    contract = json.loads(paths["contract_json"].read_text(encoding="utf-8"))
    diff = json.loads(paths["diff_json"].read_text(encoding="utf-8"))
    assert contract["source_sha256"] == before
    assert contract["mode"] == "shadow"
    assert contract["keypress_count"] == 0
    assert diff["stable_id_unique"] is True
    assert "\ufffd" not in paths["contract_json"].read_text(encoding="utf-8")

    with pytest.raises(ValueError, match="源文件所在目录"):
        write_question_contract_v2_shadow(source, source_dir)
