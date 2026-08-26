from __future__ import annotations

import ast
import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from shared_core.models import DocNode, QuestionUnit
from shared_core.question_understanding_v2_shadow import (
    build_question_understanding_v2_shadow,
    build_rule_question_understanding_proposal,
    write_question_understanding_v2_shadow,
)


PROJECT_ROOT = Path(__file__).resolve().parent
DOCUMENT_SHA256 = hashlib.sha256("题目理解影子".encode("utf-8")).hexdigest()


def _unit(
    question_id: str,
    source_span: tuple[int, int],
    *,
    overlay: str | None = None,
    material_blocks: list[str] | None = None,
    option_blocks: list[str] | None = None,
    warnings: list[str] | None = None,
) -> QuestionUnit:
    return QuestionUnit(
        question_id=question_id,
        subject="文科",
        subject_overlay=overlay,
        grade_hint="高三",
        question_type="choice" if option_blocks else "subjective",
        stem_blocks=[f"{question_id}．题干"],
        option_blocks=option_blocks or [],
        subquestions=[],
        media_blocks=[],
        material_blocks=material_blocks or [],
        source_span=source_span,
        confidence=0.9,
        warnings=warnings or [],
    )


def test_rule_candidate_excludes_known_internal_heading_but_keeps_surrounding_refs() -> None:
    nodes = [
        DocNode(1, "材料正文", metadata={"source_paragraph_index": 1}),
        DocNode(2, "（一）高考专练", metadata={"source_paragraph_index": 2}),
        DocNode(3, "1．题干", metadata={"source_paragraph_index": 3}),
    ]
    legacy = _unit(
        "1",
        (1, 3),
        overlay="zhongmei_chinese",
        material_blocks=["材料正文"],
    )

    proposal, refs = build_rule_question_understanding_proposal(
        DOCUMENT_SHA256,
        (legacy,),
        nodes,
    )

    decisions = proposal.units[0].ref_decisions
    assert [(item.disposition, item.role) for item in decisions] == [
        ("include", "material"),
        ("exclude", "internal_heading"),
        ("include", "stem"),
    ]
    assert len(refs) == 3


def test_rule_candidate_classifies_each_media_hash_independently() -> None:
    decorative_hash = (
        "363013449da1db958e2f66f789d87f7eede936259833c21c1b645b8dc3953e71"
    )
    question_hash = hashlib.sha256(b"question-media").hexdigest()
    nodes = [
        DocNode(
            1,
            "1．根据图示作答",
            has_inline_media=True,
            metadata={
                "source_paragraph_index": 1,
                "media_sha256": [decorative_hash, question_hash],
            },
        )
    ]
    legacy = _unit("1", (1, 1), overlay="future_biology")

    proposal, _ = build_rule_question_understanding_proposal(
        DOCUMENT_SHA256,
        (legacy,),
        nodes,
    )

    media_decisions = [
        item
        for item in proposal.units[0].ref_decisions
        if item.source_ref.native_kind == "media"
    ]
    assert [(item.disposition, item.role) for item in media_decisions] == [
        ("exclude", "decorative_media"),
        ("include", "question_media"),
    ]


def test_shadow_diff_compares_each_ref_and_keeps_old_production_disabled() -> None:
    nodes = [
        DocNode(1, "1．第一题", metadata={"source_paragraph_index": 1}),
        DocNode(2, "2．第二题", metadata={"source_paragraph_index": 2}),
    ]
    legacy_units = (_unit("1", (1, 1)), _unit("2", (2, 2)))

    bundle = build_question_understanding_v2_shadow(
        DOCUMENT_SHA256,
        legacy_units,
        nodes,
    )

    result = bundle["result"]
    diff = bundle["diff"]
    assert result["status"] == "compiled_for_review"
    assert result["unit_count"] == 2
    assert result["production_execution_enabled"] is False
    assert result["keypress_count"] == 0
    assert diff["legacy_unit_count"] == diff["v2_unit_count"] == 2
    assert all(item["ref_decisions"] for item in diff["comparisons"])
    assert all(
        {decision["disposition"] for decision in item["ref_decisions"]}
        == {"include"}
        for item in diff["comparisons"]
    )


def test_shadow_writer_is_readonly_utf8_and_requires_external_output_dir(
    tmp_path: Path,
) -> None:
    source_dir = tmp_path / "source"
    output_dir = tmp_path / "artifacts"
    source_dir.mkdir()
    source = source_dir / "普通连续题.docx"
    document = Document()
    document.add_paragraph("1．第一题")
    document.add_paragraph("2．第二题")
    document.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()

    paths = write_question_understanding_v2_shadow(source, output_dir)

    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert set(paths) == {"candidates_json", "question_units_json", "diff_json"}
    for path in paths.values():
        content = path.read_text(encoding="utf-8")
        assert "\ufffd" not in content
        json.loads(content)
    result = json.loads(paths["question_units_json"].read_text(encoding="utf-8"))
    assert result["mode"] == "shadow"
    assert result["production_execution_enabled"] is False
    assert result["keypress_count"] == 0
    assert result["unit_count"] == 2

    with pytest.raises(ValueError, match="源文件所在目录"):
        write_question_understanding_v2_shadow(source, source_dir)


def test_shadow_modules_do_not_import_execution_renderer_or_answer_dependencies() -> None:
    forbidden = {
        "pyautogui",
        "win32com",
        "wps_helper",
        "answer_input",
        "core_parser",
        "document_render",
    }
    for relative_path in (
        "shared_core/question_understanding_v2.py",
        "shared_core/question_understanding_v2_shadow.py",
        "tools/build_question_units_v2_shadow.py",
    ):
        path = PROJECT_ROOT / relative_path
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        imported: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.update(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.add(node.module)
        assert not {
            name
            for name in imported
            if any(part in forbidden for part in name.split("."))
        }, relative_path
