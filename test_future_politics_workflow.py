from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.shared import Cm

from shared_core import (
    build_answer_units_from_docx,
    detect_subject_overlay,
    is_question_input_excluded_for_context,
)
from tools.process_future_politics import (
    DOCUMENT_XML,
    PAGE_MARGIN_TWIPS,
    W,
    _document_body,
    _element_text,
    _page_margins_twips,
    _parse_xml,
    clean_answer_unit,
    discover_units,
    split_question_document,
)


def _set_margins(document: Document, value: float) -> None:
    for section in document.sections:
        section.top_margin = Cm(value)
        section.bottom_margin = Cm(value)
        section.left_margin = Cm(value)
        section.right_margin = Cm(value)


def _write_question_source(path: Path) -> None:
    document = Document()
    _set_margins(document, 2)
    document.add_paragraph("第一单元")
    document.add_paragraph("第一课 国体与政体")
    document.add_paragraph("第一框 国家是什么")
    document.add_paragraph("基础过关练")
    document.add_paragraph("题组一 国家主权")
    document.add_paragraph("1.第一题")
    document.add_paragraph("A.甲 B.乙")
    document.add_page_break()
    document.add_paragraph("题组二 国体与政体")
    document.add_paragraph("2.第二题")
    document.add_paragraph("A.甲 B.乙")
    document.add_page_break()
    document.add_paragraph("第二课 国家的结构形式")
    document.add_paragraph("第二框 单一制和复合制")
    document.add_paragraph("易错点 1 国家的属性")
    document.add_paragraph("1.第三题")
    document.add_paragraph("A.甲 B.乙")
    document.save(path)


def _write_answer_source(path: Path) -> None:
    document = Document()
    _set_margins(document, 2.4)
    document.add_heading("第一框 国家是什么", level=1)
    document.add_heading("基础过关练", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "1.A"
    table.cell(0, 1).text = "2.D"
    document.add_paragraph("1.A 第一题解析。")
    document.add_paragraph("续表")
    document.add_heading("国导师点睛", level=1)
    document.add_paragraph("这段知识卡片必须忽略。")
    document.add_heading("日易混辨析", level=1)
    document.add_paragraph("这段易混辨析也必须忽略。")
    document.add_paragraph("2.D 第二题解析。")
    document.add_paragraph("见续表说明，正文不能误删。")
    document.add_heading("第二框 单一制和复合制", level=1)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "1.B"
    document.add_paragraph("1.B 第三题解析。")
    document.save(path)


class FuturePoliticsQuestionTests(unittest.TestCase):
    def test_overlay_ignores_prelude_group_and_error_headings(self):
        self.assertEqual(
            detect_subject_overlay("未来-高二-政治", "", base_subject="文科"),
            "future_politics",
        )
        for text in (
            "基础过关练",
            "题组一 国家主权",
            "题组三 国际关系",
            "易错点 1 国家的属性",
            "易错点4 一国两制方针",
        ):
            self.assertTrue(
                is_question_input_excluded_for_context(text, "future_politics"),
                text,
            )

    def test_split_preserves_exact_body_slice_and_two_cm_margins(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir) / "未来-高二-政治"
            root.mkdir()
            source = root / "高中政治选必1.docx"
            _write_question_source(source)
            specs = discover_units(source)
            self.assertEqual(
                [spec.output_title for spec in specs],
                ["第一框 国家是什么", "第二框 单一制和复合制"],
            )
            results = split_question_document(
                source,
                root / "题目（已拆分）",
                expected_unit_count=2,
            )
            self.assertEqual(len(results), 2)
            with ZipFile(results[0]["path"]) as package:
                written = _parse_xml(package, DOCUMENT_XML)
            self.assertEqual(
                _page_margins_twips(written),
                (PAGE_MARGIN_TWIPS,) * 4,
            )
            texts = [
                _element_text(child).strip()
                for child in _document_body(written)
                if child.tag.endswith("}p")
            ]
            self.assertEqual(texts[0], "第一框 国家是什么")
            self.assertNotIn("第二课 国家的结构形式", texts)
            self.assertEqual(
                len(
                    written.xpath(
                        './/w:br[@w:type="page"]',
                        namespaces={"w": W[1:-1]},
                    )
                ),
                1,
            )
            self.assertEqual(results[0]["removed_trailing_page_breaks"], 1)


class FuturePoliticsAnswerTests(unittest.TestCase):
    def test_cleaner_removes_cards_and_writes_reviewable_answer_blocks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir) / "未来-高二-政治"
            root.mkdir()
            question_source = root / "高中政治选必1.docx"
            _write_question_source(question_source)
            question_results = split_question_document(
                question_source,
                root / "题目（已拆分）",
                expected_unit_count=2,
            )
            answer_source = root / "高中政治选必1答案.docx"
            _write_answer_source(answer_source)
            answer_specs = discover_units(answer_source)
            with ZipFile(answer_source) as package:
                answer_root = _parse_xml(package, DOCUMENT_XML)
            result = clean_answer_unit(
                answer_source,
                answer_root,
                answer_specs[0],
                question_results[0],
                root / "答案" / "按最小单元拆分",
            )
            output = result["path"]
            document = Document(output)
            all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertNotIn("国导师点睛", all_text)
            self.assertNotIn("日易混辨析", all_text)
            self.assertNotIn("知识卡片", all_text)
            self.assertNotIn("\n续表\n", f"\n{all_text}\n")
            self.assertIn("见续表说明，正文不能误删。", all_text)
            self.assertIn("第一题解析", all_text)
            self.assertIn("第二题解析", all_text)
            self.assertEqual(
                [unit.question_id for unit in build_answer_units_from_docx(output)],
                ["1", "2"],
            )
            self.assertEqual(
                [round(value.cm, 2) for value in (
                    document.sections[0].top_margin,
                    document.sections[0].bottom_margin,
                    document.sections[0].left_margin,
                    document.sections[0].right_margin,
                )],
                [2.0, 2.0, 2.0, 2.0],
            )
            self.assertTrue(getattr(result["path"], "is_file")())


if __name__ == "__main__":
    unittest.main()
