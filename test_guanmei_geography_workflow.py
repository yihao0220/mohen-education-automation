from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.shared import Cm
from lxml import etree as ET

from shared_core import build_answer_units_from_docx, get_review_gate_result
from tools.process_guanmei_geography import (
    PAGE_MARGIN_TWIPS,
    _page_margins_twips,
    _slice_paragraph,
    clean_answer_document,
    discover_sections,
    split_question_document,
)


TINY_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
    b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
    b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xff\xff?\x00\x05"
    b"\xfe\x02\xfeA\xe2)\xb7\x00\x00\x00\x00IEND\xaeB`\x82"
)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _add_question_section(
    document: Document,
    heading: str,
    *,
    question_count: int,
    with_picture: bool = False,
) -> None:
    document.add_heading(heading, level=1)
    for qid in range(1, question_count + 1):
        paragraph = document.add_paragraph(f"{qid}．第{qid}题")
        if with_picture and qid == 1:
            paragraph.add_run().add_picture(BytesIO(TINY_PNG))


def _write_question_source(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    _add_question_section(
        document,
        "课时分层作业(一) 测试",
        question_count=2,
        with_picture=True,
    )
    _add_question_section(document, "微点拓展专练(一) 测试", question_count=1)
    _add_question_section(document, "章末综合测评(一) 测试", question_count=2)
    document.save(path)


def _append_omath(paragraph, value: str) -> None:
    math = ET.fromstring(
        (
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/'
            'officeDocument/2006/math"><m:r><m:t>'
            f"{value}"
            "</m:t></m:r></m:oMath>"
        ).encode("utf-8")
    )
    paragraph._p.append(math)


def _write_answer_source(path: Path) -> None:
    document = Document()
    document.add_paragraph("课时分层作业(一)")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "1"
    table.cell(1, 0).text = "A"
    document.add_paragraph("2．(1)主观题答案")
    picture = document.add_paragraph()
    picture.add_run().add_picture(BytesIO(TINY_PNG))
    math_paragraph = document.add_paragraph("(2)")
    _append_omath(math_paragraph, "x=1")
    document.add_paragraph("试题精析")
    document.add_paragraph(
        "1．A [第1题，选择题解析。]"
    )
    document.add_paragraph("2．第(1)题，主观题解析。")
    document.save(path)


class GuanmeiGeographyQuestionSplitTests(unittest.TestCase):
    def test_split_preserves_body_picture_and_exact_two_cm_margins(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "新高二课时分层作业.docx"
            _write_question_source(source)
            source_hash = _digest(source)

            specs = discover_sections(source)
            self.assertEqual(
                [(spec.kind, spec.number) for spec in specs],
                [
                    ("课时分层作业", 1),
                    ("微点拓展专练", 1),
                    ("章末综合测评", 1),
                ],
            )
            results = split_question_document(
                source,
                root,
                expected_counts=None,
            )

            self.assertEqual(_digest(source), source_hash)
            self.assertEqual(len(results), 3)
            lesson = Path(results[0].output)
            with ZipFile(lesson) as package:
                xml_root = ET.fromstring(package.read("word/document.xml"))
                self.assertEqual(
                    len(
                        xml_root.findall(
                            ".//w:drawing",
                            namespaces={
                                "w": (
                                    "http://schemas.openxmlformats.org/"
                                    "wordprocessingml/2006/main"
                                )
                            },
                        )
                    ),
                    1,
                )
            self.assertEqual(
                _page_margins_twips(xml_root),
                (PAGE_MARGIN_TWIPS,) * 4,
            )
            self.assertEqual(results[0].question_count, 2)


class GuanmeiGeographyAnswerCleanTests(unittest.TestCase):
    def test_rich_paragraph_slice_does_not_duplicate_out_of_range_math(self):
        document = Document()
        paragraph = document.add_paragraph("左段")
        _append_omath(paragraph, "x=1")
        paragraph.add_run("右段")
        sliced = _slice_paragraph(paragraph._p, 5, 7)

        self.assertEqual(
            "".join(
                node.text or ""
                for node in sliced.iter()
                if ET.QName(node).localname in {"t"}
            ),
            "右段",
        )
        self.assertEqual(
            sum(1 for node in sliced.iter() if ET.QName(node).localname == "oMath"),
            0,
        )

    def test_clean_answer_preserves_question_media_math_and_review_gate(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "课时分层作业1　参考答案.docx"
            output = root / "课时分层作业1　参考答案_已清洗.docx"
            _write_answer_source(source)
            source_hash = _digest(source)

            result = clean_answer_document(
                source,
                output,
                expected_question_ids=["1", "2"],
                kind="课时分层作业",
                number=1,
            )

            self.assertEqual(_digest(source), source_hash)
            self.assertEqual(result.answer_count, 2)
            self.assertEqual(result.placeholder_analysis_count, 0)
            self.assertGreaterEqual(result.picture_count, 1)
            self.assertGreaterEqual(result.math_count, 1)
            units = build_answer_units_from_docx(
                output,
                preserve_source_positions=True,
            )
            self.assertEqual([unit.question_id for unit in units], ["1", "2"])
            self.assertEqual(units[0].answer_items[0].text, "A")
            self.assertIn("选择题解析", units[0].analysis_items[0].text)
            self.assertTrue(get_review_gate_result(output)["allowed"])


if __name__ == "__main__":
    unittest.main()
