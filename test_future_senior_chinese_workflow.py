from __future__ import annotations

from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest

from docx import Document
from docx.shared import Cm

sys.modules.setdefault("pyautogui", SimpleNamespace(press=lambda *_args, **_kwargs: None))
sys.modules.setdefault("wps_helper", SimpleNamespace(get_active_wps=lambda: None))

from shared_core import build_answer_units_from_docx
from 答案录入.answer_input import (
    build_blocks_from_units,
    find_subquestion_matches,
    should_split_subquestion_answers,
    strip_input_question_prefix,
)
from tools.process_future_senior_chinese import (
    clean_answer_section,
    discover_answer_sections,
    discover_question_parts,
    match_answer_sections,
    normalize_exam_title,
    inspect_question_document,
    split_question_document,
)


def _set_margins(document: Document, value: float = 2.0) -> None:
    for section in document.sections:
        section.top_margin = Cm(value)
        section.bottom_margin = Cm(value)
        section.left_margin = Cm(value)
        section.right_margin = Cm(value)


def _write_question_source(path: Path) -> None:
    document = Document()
    _set_margins(document)
    document.add_heading("现代文阅读真题精炼", level=1)
    document.add_heading("2021新高考I卷", level=1)
    document.add_paragraph("阅读下面的文字，完成1~2题。")
    document.add_paragraph("1．第一题")
    document.add_paragraph("2．第二题")
    document.add_heading("2021新高考Ⅱ卷", level=1)
    document.add_paragraph("阅读下面的文字，完成1~2题。")
    document.add_paragraph("1．第三题")
    document.add_paragraph("2．第四题")
    document.add_heading("语言文字运用真题精炼", level=1)
    document.add_heading("2021新高考I卷", level=1)
    document.add_paragraph("阅读下面的文言文，完成10~11题。")
    document.add_paragraph("10．第五题")
    document.add_paragraph("11．第六题")
    document.add_heading("2021新高考Ⅱ卷", level=1)
    document.add_paragraph("阅读下面的文言文，完成18~19题。")
    document.add_paragraph("18．第七题")
    document.add_paragraph("19．第八题")
    document.save(path)


def _write_answer_source(path: Path) -> None:
    document = Document()
    _set_margins(document)
    document.add_heading("2021 新高考 I 卷", level=1)
    document.add_paragraph("1. A")
    document.add_paragraph("【解析】第一题解析")
    document.add_paragraph("2. 第二题答案")
    document.add_paragraph("【解析】第二题解析")
    document.add_heading("2021 新高考 II 卷", level=1)
    document.add_paragraph("1. B")
    document.add_paragraph("【解析】第三题解析")
    document.add_paragraph("2. D")
    document.add_paragraph("【解析】第四题解析")
    document.add_heading("2021 新高考 | 卷", level=1)
    document.add_paragraph("10. C")
    document.add_paragraph("【解析】第五题解析")
    document.add_paragraph("11. 第六题答案")
    document.add_paragraph("【解析】第六题解析")
    document.add_heading("2021 新高考 II 卷", level=1)
    document.add_paragraph("18. A")
    document.add_paragraph("【解析】第七题解析")
    document.add_paragraph("19. B")
    document.add_paragraph("【解析】第八题解析")
    document.save(path)


class FutureSeniorChineseQuestionTests(unittest.TestCase):
    def test_directory_order_drives_prefix_and_first_part_keeps_category(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "questions.docx"
            _write_question_source(source)
            specs = discover_question_parts(source, expected_count=4)
            self.assertEqual(
                [spec.filename for spec in specs],
                [
                    "01 2021新高考I卷.docx",
                    "02 2021新高考Ⅱ卷.docx",
                    "03 2021新高考I卷.docx",
                    "04 2021新高考Ⅱ卷.docx",
                ],
            )
            results = split_question_document(
                source,
                root / "题目（已拆分）",
                expected_count=4,
            )
            first = Document(results[0]["path"])
            third = Document(results[2]["path"])
            self.assertEqual(first.paragraphs[0].text, "现代文阅读真题精炼")
            self.assertEqual(first.paragraphs[1].text, "2021新高考I卷")
            self.assertEqual(third.paragraphs[0].text, "语言文字运用真题精炼")
            self.assertEqual(third.paragraphs[1].text, "2021新高考I卷")
            self.assertEqual(
                [round(value.cm, 2) for value in (
                    first.sections[0].top_margin,
                    first.sections[0].bottom_margin,
                    first.sections[0].left_margin,
                    first.sections[0].right_margin,
                )],
                [2.0, 2.0, 2.0, 2.0],
            )


class FutureSeniorChineseAnswerTests(unittest.TestCase):
    def test_title_normalization_handles_roman_glyphs_and_bars(self):
        expected = "2021新高考II卷"
        for value in (
            "2021 新高考 II 卷",
            "2021新高考Ⅱ卷",
            "2021 新高考 || 卷",
        ):
            self.assertEqual(normalize_exam_title(value), expected)

    def test_answers_follow_question_prefix_and_restart_from_one(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            question_source = root / "questions.docx"
            answer_source = root / "answers.docx"
            _write_question_source(question_source)
            _write_answer_source(answer_source)
            question_results = split_question_document(
                question_source,
                root / "题目（已拆分）",
                expected_count=4,
            )
            matches = match_answer_sections(
                question_results,
                discover_answer_sections(answer_source),
            )
            self.assertTrue(all(item["status"] == "matched" for item in matches))
            outputs = [
                clean_answer_section(answer_source, root / "答案", item)
                for item in matches
            ]
            self.assertEqual(
                [item["path"].name for item in outputs],
                [
                    "01 2021新高考I卷-答案_已清洗.docx",
                    "02 2021新高考Ⅱ卷-答案_已清洗.docx",
                    "03 2021新高考I卷-答案_已清洗.docx",
                    "04 2021新高考Ⅱ卷-答案_已清洗.docx",
                ],
            )
            for output in outputs:
                units = build_answer_units_from_docx(output["path"])
                self.assertEqual([unit.question_id for unit in units], ["1"])
                self.assertEqual(units[0].answer_mode, "subquestion")
                self.assertEqual(
                    [item.item_id for item in units[0].answer_items],
                    ["（1）", "（2）"],
                )
                block = build_blocks_from_units(units)[0]
                paragraphs = [p.text for p in Document(output["path"]).paragraphs]
                answer_text = "\n".join(
                    paragraphs[
                        block["ans_start_p"] - 1 : block["ana_start_p"] - 1
                    ]
                )
                _, text_body = strip_input_question_prefix(answer_text)
                self.assertEqual(len(find_subquestion_matches(text_body)), 2)
                self.assertTrue(
                    should_split_subquestion_answers(SimpleNamespace(Name="答案.docx"), block)
                )
                text = "\n".join(p.text for p in Document(output["path"]).paragraphs)
                self.assertIn("解析：", text)
                self.assertNotIn("【解析】", text)

    def test_two_reading_groups_restart_subquestion_numbers(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            question_source = root / "questions.docx"
            answer_source = root / "answers.docx"

            question = Document()
            _set_margins(question)
            question.add_heading("现代文阅读真题精炼", level=1)
            question.add_heading("2019全国新课标I卷", level=1)
            question.add_paragraph("阅读下面的文字，完成1~3题。")
            for number in range(1, 4):
                question.add_paragraph(f"{number}．第一组第{number}题")
            question.add_paragraph("阅读下面的文字，完成4~6题。")
            for number in range(4, 7):
                question.add_paragraph(f"{number}．第二组第{number}题")
            question.save(question_source)

            answer = Document()
            _set_margins(answer)
            answer.add_heading("2019 全国新课标 I 卷", level=1)
            for number, value in enumerate("ABCDEF", 1):
                answer.add_paragraph(f"{number}. {value}")
                answer.add_paragraph(f"【解析】第{number}题解析")
            answer.save(answer_source)

            question_results = split_question_document(
                question_source,
                root / "题目（已拆分）",
                expected_count=1,
            )
            self.assertEqual(
                question_results[0]["question_groups"],
                (("1", "2", "3"), ("4", "5", "6")),
            )
            matches = match_answer_sections(
                question_results,
                discover_answer_sections(answer_source),
            )
            output = clean_answer_section(answer_source, root / "答案", matches[0])
            units = build_answer_units_from_docx(output["path"])

            self.assertEqual([unit.question_id for unit in units], ["1", "2"])
            self.assertEqual([unit.answer_mode for unit in units], ["subquestion"] * 2)
            self.assertEqual([len(unit.answer_items) for unit in units], [3, 3])
            self.assertEqual(output["f4_answer_count"], 6)
            paragraphs = [p.text for p in Document(output["path"]).paragraphs]
            self.assertEqual(paragraphs.count("（1）A") + paragraphs.count("（1）D"), 2)
            self.assertEqual(paragraphs.count("解析："), 2)
            for block in build_blocks_from_units(units):
                answer_text = "\n".join(
                    paragraphs[
                        block["ans_start_p"] - 1 : block["ana_start_p"] - 1
                    ]
                )
                _, text_body = strip_input_question_prefix(answer_text)
                self.assertEqual(len(find_subquestion_matches(text_body)), 3)
                self.assertTrue(
                    should_split_subquestion_answers(SimpleNamespace(Name="答案.docx"), block)
                )

    def test_question_answer_id_mismatch_is_skipped(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            question_source = root / "questions.docx"
            answer_source = root / "answers.docx"
            _write_question_source(question_source)
            _write_answer_source(answer_source)
            document = Document(answer_source)
            document.add_paragraph("20. 多余答案")
            document.add_paragraph("【解析】多余解析")
            document.save(answer_source)
            question_results = split_question_document(
                question_source,
                root / "题目（已拆分）",
                expected_count=4,
            )
            matches = match_answer_sections(
                question_results,
                discover_answer_sections(answer_source),
            )
            self.assertEqual(matches[-1]["status"], "skipped")
            self.assertIn("题目原题号为 18/19", matches[-1]["reason"])

    def test_category_heading_does_not_block_matching(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            question_source = root / "questions.docx"
            answer_source = root / "answers.docx"
            _write_question_source(question_source)
            _write_answer_source(answer_source)
            question_results = split_question_document(
                question_source,
                root / "题目（已拆分）",
                expected_count=4,
            )
            matches = match_answer_sections(
                question_results,
                discover_answer_sections(answer_source),
            )
            self.assertEqual(matches[2]["status"], "matched")

    def test_missing_middle_answer_does_not_block_later_matches(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            question_source = root / "questions.docx"
            answer_source = root / "answers.docx"
            _write_question_source(question_source)
            document = Document()
            _set_margins(document)
            document.add_heading("2021 新高考 I 卷", level=1)
            document.add_paragraph("1. A")
            document.add_paragraph("【解析】第一题")
            document.add_paragraph("2. B")
            document.add_paragraph("【解析】第二题")
            document.add_heading("2021 新高考 | 卷", level=1)
            document.add_paragraph("10. C")
            document.add_paragraph("【解析】第五题")
            document.add_paragraph("11. D")
            document.add_paragraph("【解析】第六题")
            document.add_heading("2021 新高考 II 卷", level=1)
            document.add_paragraph("18. A")
            document.add_paragraph("【解析】第七题")
            document.add_paragraph("19. B")
            document.add_paragraph("【解析】第八题")
            document.save(answer_source)
            question_results = inspect_question_document(
                question_source,
                expected_count=4,
            )
            matches = match_answer_sections(
                question_results,
                discover_answer_sections(answer_source),
            )
            self.assertEqual(
                [item["status"] for item in matches],
                ["matched", "skipped", "matched", "matched"],
            )


if __name__ == "__main__":
    unittest.main()
