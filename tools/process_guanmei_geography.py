from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import asdict, dataclass
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path, PurePosixPath
import re
import subprocess
import sys
import tarfile
import tempfile
import unicodedata
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from lxml import etree as ET


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from shared_core import (  # noqa: E402
    build_answer_units_from_docx,
    build_question_units_from_docx,
    build_review_report,
    get_review_gate_result,
    initialize_review_status,
    map_answers,
    update_review_status,
)
from shared_core.answer_core import infer_grouped_question_ids  # noqa: E402


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
IMAGE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
)
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
W = f"{{{W_NS}}}"
M = f"{{{M_NS}}}"
R = f"{{{R_NS}}}"
PKG_REL = f"{{{PKG_REL_NS}}}"

DOCUMENT_XML = "word/document.xml"
DOCUMENT_RELS = "word/_rels/document.xml.rels"
PAGE_MARGIN_TWIPS = 1134

SECTION_PATTERN = re.compile(
    r"^(?P<kind>课时分层作业|微点拓展专练|章末综合测评)"
    r"[\(（](?P<number>[一二三四五六七八九十百0-9]+)[\)）]"
)
QUESTION_ID_PATTERN = re.compile(r"^\s*(?P<qid>\d+)\s*[．.]")
ANALYSIS_ITEM_PATTERN = re.compile(r"第(?P<qid>\d+)题[，,]")
SUBJECTIVE_ANALYSIS_PATTERN = re.compile(r"^\s*(?P<qid>\d+)\s*[．.]")
DELIVERY_EXCLUDED_PARTS = {".DS_Store"}

EXPECTED_SECTION_COUNTS = {
    "课时分层作业": 24,
    "微点拓展专练": 8,
    "章末综合测评": 5,
}


@dataclass(frozen=True)
class SectionSpec:
    kind: str
    number: int
    heading: str
    start: int
    end: int

    @property
    def output_name(self) -> str:
        return f"{self.kind}{self.number}.docx"


@dataclass(frozen=True)
class DocumentResult:
    source: str
    output: str
    kind: str
    number: int
    question_count: int
    page_margins_twips: tuple[int, int, int, int]
    sha256: str


@dataclass(frozen=True)
class AnswerResult:
    source: str
    output: str
    kind: str
    number: int
    answer_count: int
    placeholder_analysis_count: int
    picture_count: int
    math_count: int
    sha256: str


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _element_text(element) -> str:
    return "".join(
        node.text or ""
        for node in element.iter()
        if node.tag in {f"{W}t", f"{M}t"}
    )


def _all_element_text(element) -> str:
    return "".join(node.text or "" for node in element.iter())


def _set_text(node, value: str) -> None:
    node.text = value
    if value.startswith(" ") or value.endswith(" "):
        node.set(XML_SPACE, "preserve")
    else:
        node.attrib.pop(XML_SPACE, None)


def _parse_xml(package: ZipFile, name: str):
    return ET.fromstring(package.read(name))


def _document_body(root):
    body = root.find(f"{W}body")
    if body is None:
        raise ValueError("DOCX 缺少 word/document.xml 正文节点")
    return body


def _parse_chinese_number(value: str) -> int:
    if value.isdigit():
        return int(value)
    digits = {
        "零": 0,
        "一": 1,
        "二": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
    }
    if value == "十":
        return 10
    if "十" in value:
        left, right = value.split("十", 1)
        return (digits.get(left, 1) * 10) + digits.get(right, 0)
    if "百" in value:
        left, right = value.split("百", 1)
        return digits.get(left, 1) * 100 + _parse_chinese_number(right or "零")
    if value not in digits:
        raise ValueError(f"无法识别中文序号: {value}")
    return digits[value]


def discover_sections(source_docx: str | Path) -> list[SectionSpec]:
    source = Path(source_docx)
    with ZipFile(source) as package:
        root = _parse_xml(package, DOCUMENT_XML)
    children = list(_document_body(root))
    headings: list[tuple[int, str, int, str]] = []
    for index, child in enumerate(children):
        if child.tag != f"{W}p":
            continue
        heading = _element_text(child).strip()
        match = SECTION_PATTERN.match(heading)
        if match:
            headings.append(
                (
                    index,
                    match.group("kind"),
                    _parse_chinese_number(match.group("number")),
                    heading,
                )
            )
    specs = []
    for position, (start, kind, number, heading) in enumerate(headings):
        end = headings[position + 1][0] if position + 1 < len(headings) else len(children)
        while end > start and children[end - 1].tag == f"{W}sectPr":
            end -= 1
        specs.append(SectionSpec(kind, number, heading, start, end))
    return specs


def validate_section_structure(
    specs: list[SectionSpec],
    expected_counts: dict[str, int] | None = EXPECTED_SECTION_COUNTS,
) -> None:
    if not specs:
        raise ValueError("总题目文档中未识别到章节标题")
    if expected_counts is None:
        return
    for kind, expected_count in expected_counts.items():
        actual = sorted(spec.number for spec in specs if spec.kind == kind)
        expected = list(range(1, expected_count + 1))
        if actual != expected:
            raise ValueError(
                f"{kind} 章节序号异常: 期望 {expected}，实际 {actual}"
            )


def _source_part_from_rels_name(rels_name: str) -> str | None:
    path = PurePosixPath(rels_name)
    if path.name == ".rels" and str(path.parent) == "_rels":
        return ""
    if path.parent.name != "_rels" or not path.name.endswith(".rels"):
        return None
    source_name = path.name[: -len(".rels")]
    return str(path.parent.parent / source_name)


def _resolve_relationship_target(rels_name: str, target: str) -> str | None:
    if "://" in target or target.startswith("#"):
        return None
    source_part = _source_part_from_rels_name(rels_name)
    if source_part is None:
        return None
    base = PurePosixPath(source_part).parent if source_part else PurePosixPath("")
    pieces: list[str] = []
    for piece in (base / target).parts:
        if piece in {"", "."}:
            continue
        if piece == "..":
            if pieces:
                pieces.pop()
            continue
        pieces.append(piece)
    return "/".join(pieces)


def _referenced_relationship_ids(document_root) -> set[str]:
    values: set[str] = set()
    for element in document_root.iter():
        for attribute, value in element.attrib.items():
            if attribute.startswith(f"{{{R_NS}}}"):
                values.add(value)
    return values


def _prepare_relationship_overrides(
    package: ZipFile,
    document_root,
) -> tuple[dict[str, bytes], set[str]]:
    overrides: dict[str, bytes] = {}
    removed_media: set[str] = set()
    if DOCUMENT_RELS not in package.namelist():
        return overrides, removed_media

    referenced_ids = _referenced_relationship_ids(document_root)
    rels_root = _parse_xml(package, DOCUMENT_RELS)
    for relationship in list(rels_root):
        if (
            relationship.get("Type") == IMAGE_REL_TYPE
            and relationship.get("Id") not in referenced_ids
        ):
            target = _resolve_relationship_target(
                DOCUMENT_RELS,
                relationship.get("Target", ""),
            )
            if target:
                removed_media.add(target)
            rels_root.remove(relationship)
    overrides[DOCUMENT_RELS] = ET.tostring(
        rels_root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    remaining_media: set[str] = set()
    for rels_name in package.namelist():
        if not rels_name.endswith(".rels"):
            continue
        rels_payload = overrides.get(rels_name, package.read(rels_name))
        rels = ET.fromstring(rels_payload)
        for relationship in rels:
            if relationship.get("Type") != IMAGE_REL_TYPE:
                continue
            target = _resolve_relationship_target(
                rels_name,
                relationship.get("Target", ""),
            )
            if target:
                remaining_media.add(target)
    return overrides, removed_media - remaining_media


def _write_docx_package(
    source_path: Path,
    output_path: Path,
    document_root,
    *,
    prune_unreferenced_images: bool = True,
) -> None:
    if source_path.resolve() == output_path.resolve():
        raise ValueError("输出文档不能覆盖源文档")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document_xml = ET.tostring(
        document_root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    with tempfile.NamedTemporaryFile(
        prefix=f"{output_path.stem}_",
        suffix=".docx",
        dir=output_path.parent,
        delete=False,
    ) as stream:
        temporary_path = Path(stream.name)
    try:
        with ZipFile(source_path) as source_package:
            overrides: dict[str, bytes] = {}
            removed_media: set[str] = set()
            if prune_unreferenced_images:
                overrides, removed_media = _prepare_relationship_overrides(
                    source_package,
                    document_root,
                )
            with ZipFile(temporary_path, "w", compression=ZIP_DEFLATED) as target:
                for info in source_package.infolist():
                    if info.filename in removed_media:
                        continue
                    payload = source_package.read(info.filename)
                    if info.filename == DOCUMENT_XML:
                        payload = document_xml
                    elif info.filename in overrides:
                        payload = overrides[info.filename]
                    target.writestr(info, payload)
        with ZipFile(temporary_path) as check:
            broken_member = check.testzip()
            if broken_member:
                raise ValueError(f"DOCX 压缩包损坏: {broken_member}")
        Document(temporary_path)
        temporary_path.replace(output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _page_margins_twips(root) -> tuple[int, int, int, int]:
    body = _document_body(root)
    sect_pr = body.find(f"{W}sectPr")
    if sect_pr is None:
        raise ValueError("文档缺少节属性，无法验证页边距")
    page_margin = sect_pr.find(f"{W}pgMar")
    if page_margin is None:
        raise ValueError("文档缺少页边距设置")
    return tuple(
        int(page_margin.get(f"{W}{name}", "-1"))
        for name in ("top", "bottom", "left", "right")
    )


def _question_ids_from_root(root) -> list[str]:
    result: list[str] = []
    for child in _document_body(root):
        if child.tag != f"{W}p":
            continue
        match = QUESTION_ID_PATTERN.match(_element_text(child))
        if match:
            result.append(match.group("qid"))
    return result


def _assert_contiguous_ids(ids: list[str], context: str) -> None:
    if not ids:
        raise ValueError(f"{context} 未识别到题号")
    actual = [int(value) for value in ids]
    expected = list(range(1, max(actual) + 1))
    if actual != expected:
        raise ValueError(f"{context} 题号不连续: 期望 {expected}，实际 {actual}")


def _canonical_children(root) -> list[bytes]:
    return [
        ET.tostring(child, method="c14n")
        for child in _document_body(root)
        if child.tag != f"{W}sectPr"
    ]


def _assert_relationships_resolve(path: Path) -> None:
    with ZipFile(path) as package:
        root = _parse_xml(package, DOCUMENT_XML)
        rels = _parse_xml(package, DOCUMENT_RELS)
        available = {relationship.get("Id") for relationship in rels}
        missing = sorted(_referenced_relationship_ids(root) - available)
    if missing:
        raise ValueError(f"{path.name} 存在失效资源关系: {missing}")


def split_question_document(
    source_docx: str | Path,
    output_root: str | Path,
    *,
    expected_counts: dict[str, int] | None = EXPECTED_SECTION_COUNTS,
) -> list[DocumentResult]:
    source = Path(source_docx)
    output_base = Path(output_root)
    specs = discover_sections(source)
    validate_section_structure(specs, expected_counts)
    with ZipFile(source) as package:
        source_root = _parse_xml(package, DOCUMENT_XML)
    source_body = _document_body(source_root)
    source_children = list(source_body)
    source_sect_pr = source_body.find(f"{W}sectPr")
    if source_sect_pr is None:
        raise ValueError("总题目文档缺少节属性")
    if _page_margins_twips(source_root) != (PAGE_MARGIN_TWIPS,) * 4:
        raise ValueError(
            "总题目文档四边页边距不是 2 cm（1134 twips），"
            "为避免擅自改版已停止拆分"
        )

    directories = {
        "课时分层作业": output_base / "题目（已拆分）",
        "章末综合测评": output_base / "题目（已拆分）",
        "微点拓展专练": output_base / "微点拓展专练（暂缺答案）",
    }
    results = []
    for spec in specs:
        selected = source_children[spec.start : spec.end]
        target_root = deepcopy(source_root)
        target_body = _document_body(target_root)
        for child in list(target_body):
            target_body.remove(child)
        for child in selected:
            target_body.append(deepcopy(child))
        target_body.append(deepcopy(source_sect_pr))

        output_path = directories[spec.kind] / spec.output_name
        _write_docx_package(source, output_path, target_root)
        with ZipFile(output_path) as package:
            written_root = _parse_xml(package, DOCUMENT_XML)
            if _canonical_children(written_root) != [
                ET.tostring(child, method="c14n") for child in selected
            ]:
                raise ValueError(f"{output_path.name} 正文 XML 与原文切片不一致")
            if _page_margins_twips(written_root) != (PAGE_MARGIN_TWIPS,) * 4:
                raise ValueError(f"{output_path.name} 四边页边距不是 2 cm")
        _assert_relationships_resolve(output_path)
        question_ids = _question_ids_from_root(written_root)
        _assert_contiguous_ids(question_ids, output_path.name)
        results.append(
            DocumentResult(
                source=str(source),
                output=str(output_path),
                kind=spec.kind,
                number=spec.number,
                question_count=len(question_ids),
                page_margins_twips=_page_margins_twips(written_root),
                sha256=_sha256(output_path),
            )
        )
    return results


def _text_nodes(paragraph) -> list:
    return [
        node
        for node in paragraph.iter()
        if node.tag in {f"{W}t", f"{M}t"}
    ]


def _trim_slice_bounds(text: str, start: int, end: int) -> tuple[int, int]:
    while start < end and text[start].isspace():
        start += 1
    while end > start and (text[end - 1].isspace() or text[end - 1] == "]"):
        end -= 1
    return start, end


def _slice_paragraph(paragraph, start: int, end: int):
    copied = deepcopy(paragraph)
    cursor = 0
    for child in list(copied):
        if child.tag == f"{W}pPr":
            continue
        nodes = _text_nodes(child)
        child_length = sum(len(node.text or "") for node in nodes)
        child_start = cursor
        child_end = cursor + child_length
        cursor = child_end
        if child_length == 0:
            if not (_has_payload(child) and start <= child_start < end):
                copied.remove(child)
            continue
        if child_end <= start or child_start >= end:
            copied.remove(child)
            continue
        node_cursor = child_start
        for node in nodes:
            value = node.text or ""
            node_start = node_cursor
            node_end = node_start + len(value)
            node_cursor = node_end
            keep_start = max(start, node_start)
            keep_end = min(end, node_end)
            if keep_start >= keep_end:
                _set_text(node, "")
            else:
                _set_text(
                    node,
                    value[keep_start - node_start : keep_end - node_start],
                )
    return copied


def _new_text_paragraph(reference, text: str):
    paragraph = ET.Element(f"{W}p")
    if reference is not None:
        properties = reference.find(f"{W}pPr")
        if properties is not None:
            paragraph.append(deepcopy(properties))
    run = ET.SubElement(paragraph, f"{W}r")
    if reference is not None:
        run_properties = reference.find(f".//{W}rPr")
        if run_properties is not None:
            run.append(deepcopy(run_properties))
    text_node = ET.SubElement(run, f"{W}t")
    _set_text(text_node, text)
    return paragraph


def _has_payload(element) -> bool:
    if _element_text(element).strip():
        return True
    payload_tags = {
        f"{W}drawing",
        f"{W}pict",
        f"{W}object",
        f"{M}oMath",
        f"{M}oMathPara",
    }
    return any(node.tag in payload_tags for node in element.iter())


def _table_choice_answers(answer_docx: Path) -> tuple[dict[str, str], dict[str, object]]:
    document = Document(answer_docx)
    choices: dict[str, str] = {}
    references: dict[str, object] = {}
    for table in document.tables:
        if len(table.rows) < 2:
            continue
        qids = [cell.text.strip() for cell in table.rows[0].cells]
        answers = [cell.text.strip() for cell in table.rows[1].cells]
        if not qids or not all(value.isdigit() for value in qids if value):
            continue
        for qid, answer, cell in zip(qids, answers, table.rows[1].cells):
            if qid and answer:
                choices[qid] = answer
                references[qid] = cell.paragraphs[0]._p
    return choices, references


def _body_banner_index(body, marker: str) -> int:
    for index, child in enumerate(body):
        if marker in _all_element_text(child):
            return index
    raise ValueError(f"答案文档中未找到“{marker}”分区")


def _extract_subjective_answers(
    body,
    start: int,
    end: int,
) -> dict[str, list]:
    answers: dict[str, list] = {}
    current_qid: str | None = None
    for child in list(body)[start:end]:
        text = _element_text(child)
        match = QUESTION_ID_PATTERN.match(text) if child.tag == f"{W}p" else None
        if match:
            current_qid = match.group("qid")
            answers.setdefault(current_qid, [])
            payload_start = match.end()
            if text[payload_start:].strip() or _has_payload(child):
                answers[current_qid].append(
                    _slice_paragraph(child, payload_start, len(text))
                )
            continue
        if current_qid is not None and _has_payload(child):
            answers[current_qid].append(deepcopy(child))
    return answers


def _extract_analyses(body, start: int, end: int) -> dict[str, list]:
    analyses: dict[str, list] = {}
    current_qid: str | None = None
    for child in list(body)[start:end]:
        if child.tag != f"{W}p":
            if current_qid is not None and _has_payload(child):
                analyses.setdefault(current_qid, []).append(deepcopy(child))
            continue
        text = _element_text(child)
        markers = list(ANALYSIS_ITEM_PATTERN.finditer(text))
        if markers:
            for index, marker in enumerate(markers):
                qid = marker.group("qid")
                segment_start = marker.end()
                segment_end = (
                    markers[index + 1].start()
                    if index + 1 < len(markers)
                    else len(text)
                )
                segment_start, segment_end = _trim_slice_bounds(
                    text,
                    segment_start,
                    segment_end,
                )
                analyses.setdefault(qid, []).append(
                    _slice_paragraph(child, segment_start, segment_end)
                )
                current_qid = qid
            continue
        subjective = SUBJECTIVE_ANALYSIS_PATTERN.match(text)
        if subjective:
            current_qid = subjective.group("qid")
            start_at = subjective.end()
            start_at, end_at = _trim_slice_bounds(text, start_at, len(text))
            analyses.setdefault(current_qid, []).append(
                _slice_paragraph(child, start_at, end_at)
            )
            continue
        if current_qid is not None and _has_payload(child):
            analyses.setdefault(current_qid, []).append(deepcopy(child))
    return analyses


def _count_xml_payload(root, local_name: str) -> int:
    return sum(1 for node in root.iter() if ET.QName(node).localname == local_name)


def clean_answer_document(
    source_docx: str | Path,
    output_docx: str | Path,
    *,
    expected_question_ids: list[str],
    kind: str,
    number: int,
) -> AnswerResult:
    source = Path(source_docx)
    output = Path(output_docx)
    choices, choice_references = _table_choice_answers(source)
    with ZipFile(source) as package:
        source_root = _parse_xml(package, DOCUMENT_XML)
    source_body = _document_body(source_root)
    source_children = list(source_body)
    analysis_banner = _body_banner_index(source_body, "试题精析")
    table_index = next(
        (
            index
            for index, child in enumerate(source_children[:analysis_banner])
            if child.tag == f"{W}tbl"
        ),
        None,
    )
    if table_index is None:
        raise ValueError(f"{source.name} 未找到答案速对表格")

    subjective_answers = _extract_subjective_answers(
        source_body,
        table_index + 1,
        analysis_banner,
    )
    analyses = _extract_analyses(
        source_body,
        analysis_banner + 1,
        len(source_children),
    )

    all_answer_ids = sorted(
        {int(value) for value in choices} | {int(value) for value in subjective_answers}
    )
    answer_ids = [str(value) for value in all_answer_ids]
    if answer_ids != expected_question_ids:
        raise ValueError(
            f"{source.name} 答案题号与题目不一致: "
            f"题目 {expected_question_ids}，答案 {answer_ids}"
        )

    fallback_reference = next(
        (
            child
            for values in subjective_answers.values()
            for child in values
            if child.tag == f"{W}p"
        ),
        None,
    )
    if fallback_reference is None:
        fallback_reference = next(
            (
                child
                for values in analyses.values()
                for child in values
                if child.tag == f"{W}p"
            ),
            None,
        )

    target_root = deepcopy(source_root)
    target_body = _document_body(target_root)
    for child in list(target_body):
        target_body.remove(child)

    placeholder_count = 0
    for qid in expected_question_ids:
        reference = choice_references.get(qid, fallback_reference)
        if qid in choices:
            target_body.append(
                _new_text_paragraph(reference, f"{qid}．{choices[qid]}")
            )
        else:
            target_body.append(_new_text_paragraph(reference, f"{qid}．"))
            for child in subjective_answers.get(qid, []):
                target_body.append(deepcopy(child))

        target_body.append(_new_text_paragraph(reference, "解析："))
        analysis_children = analyses.get(qid, [])
        if analysis_children:
            for child in analysis_children:
                target_body.append(deepcopy(child))
        else:
            placeholder_count += 1

    sect_pr = source_body.find(f"{W}sectPr")
    if sect_pr is not None:
        target_body.append(deepcopy(sect_pr))

    _write_docx_package(source, output, target_root)
    _assert_relationships_resolve(output)
    answer_units = build_answer_units_from_docx(
        output,
        preserve_source_positions=True,
    )
    parsed_ids = [unit.question_id for unit in answer_units]
    if parsed_ids != expected_question_ids:
        raise ValueError(
            f"{output.name} 清洗后解析题号异常: "
            f"期望 {expected_question_ids}，实际 {parsed_ids}"
        )
    review_flags = [
        (unit.question_id, list(unit.review_flags))
        for unit in answer_units
        if unit.review_flags
    ]
    if review_flags:
        raise ValueError(f"{output.name} 清洗后仍有审核标记: {review_flags}")

    with ZipFile(output) as package:
        output_root = _parse_xml(package, DOCUMENT_XML)
    initialize_review_status(output)
    update_review_status(
        output,
        status="approved",
        reviewer="Codex offline validation",
        note=(
            "题号、答案单元、解析占位、DOCX 压缩结构及资源关系自动检查通过；"
            "Windows WPS 生产环境仍需人工打开确认。"
        ),
    )
    gate = get_review_gate_result(output)
    if not gate["allowed"]:
        raise ValueError(f"{output.name} 审核门禁未放行: {gate['reason']}")

    return AnswerResult(
        source=str(source),
        output=str(output),
        kind=kind,
        number=number,
        answer_count=len(answer_units),
        placeholder_analysis_count=placeholder_count,
        picture_count=_count_xml_payload(output_root, "pict")
        + _count_xml_payload(output_root, "drawing"),
        math_count=_count_xml_payload(output_root, "oMath"),
        sha256=_sha256(output),
    )


def _find_answer_source(project_root: Path, kind: str, number: int) -> Path:
    directory = {
        "课时分层作业": project_root / "课时分层作业参考答案",
        "章末综合测评": project_root / "章末综合测评参考答案",
    }[kind]
    filename_pattern = re.compile(
        rf"^{re.escape(kind)}{number}(?:\s|　)*参考答案\.docx$"
    )
    matches = sorted(
        path
        for path in directory.glob("*.docx")
        if filename_pattern.match(path.name)
    )
    if len(matches) != 1:
        raise ValueError(
            f"{kind}{number} 参考答案匹配数量应为 1，实际为 {len(matches)}: "
            f"{[path.name for path in matches]}"
        )
    return matches[0]


def clean_answer_batch(
    project_root: str | Path,
    question_results: list[DocumentResult],
) -> list[AnswerResult]:
    root = Path(project_root)
    answer_output = root / "答案（已清洗）"
    results = []
    for question in question_results:
        if question.kind == "微点拓展专练":
            continue
        question_path = Path(question.output)
        with ZipFile(question_path) as package:
            question_root = _parse_xml(package, DOCUMENT_XML)
        question_ids = _question_ids_from_root(question_root)
        source = _find_answer_source(root, question.kind, question.number)
        output = answer_output / f"{source.stem}_已清洗.docx"
        results.append(
            clean_answer_document(
                source,
                output,
                expected_question_ids=question_ids,
                kind=question.kind,
                number=question.number,
            )
        )
    return results


def validate_offline_input_mapping(
    question_results: list[DocumentResult],
    answer_results: list[AnswerResult],
) -> dict:
    answers_by_section = {
        (result.kind, result.number): result for result in answer_results
    }
    action_counts = {"F1": 0, "F2": 0, "F4": 0, "F3": 0}
    document_summaries = []
    for question_result in question_results:
        if question_result.kind == "微点拓展专练":
            continue
        answer_result = answers_by_section.get(
            (question_result.kind, question_result.number)
        )
        if answer_result is None:
            raise ValueError(
                f"{question_result.kind}{question_result.number} 缺少清洗答案"
            )
        question_units = build_question_units_from_docx(
            question_result.output,
            grade_hint="高二",
        )
        answer_units = build_answer_units_from_docx(
            answer_result.output,
            preserve_source_positions=True,
        )
        covered_question_ids = [
            qid
            for unit in question_units
            for qid in infer_grouped_question_ids(unit)
        ]
        answer_ids = [unit.question_id for unit in answer_units]
        if covered_question_ids != answer_ids:
            raise ValueError(
                f"{Path(question_result.output).name} F1 题号覆盖与答案不一致: "
                f"题目 {covered_question_ids}，答案 {answer_ids}"
            )

        mapped = map_answers(question_units, answer_units)
        risky_mappings = [
            (unit.question_id, list(unit.review_flags))
            for unit in mapped
            if unit.review_flags
        ]
        report = build_review_report(
            Path(question_result.output).name,
            question_units,
            mapped,
        )
        if (
            len(mapped) != len(question_units)
            or risky_mappings
            or report.summary["high_risk_count"]
        ):
            raise ValueError(
                f"{Path(question_result.output).name} 离线题答映射未通过: "
                f"F1={len(question_units)}，映射={len(mapped)}，"
                f"风险={risky_mappings}，"
                f"高风险数={report.summary['high_risk_count']}"
            )

        local_counts = {
            "F1": len(question_units),
            "F2": sum(unit.answer_mode != "subquestion" for unit in answer_units),
            "F4": sum(
                len(unit.answer_items)
                for unit in answer_units
                if unit.answer_mode == "subquestion"
            ),
            "F3": len(answer_units),
        }
        for key, value in local_counts.items():
            action_counts[key] += value
        document_summaries.append(
            {
                "question": Path(question_result.output).name,
                "answer": Path(answer_result.output).name,
                "logical_question_ids": len(answer_ids),
                "f1_blocks": len(question_units),
                "mapped_blocks": len(mapped),
                "high_risk_count": report.summary["high_risk_count"],
                "actions": local_counts,
            }
        )
    return {
        "mode": "offline_no_key_rehearsal",
        "keypress_execution_enabled": False,
        "document_pairs": len(document_summaries),
        "logical_question_ids_covered": sum(
            item["logical_question_ids"] for item in document_summaries
        ),
        "action_counts": action_counts,
        "mapping_errors": 0,
        "documents": document_summaries,
        "windows_wps_plugin_execution": "pending",
    }


def _normalized_zip_name(value: str) -> str:
    return unicodedata.normalize("NFC", value.replace("\\", "/"))


def _zip_info(name: str) -> ZipInfo:
    info = ZipInfo(_normalized_zip_name(name))
    info.date_time = (2026, 1, 1, 0, 0, 0)
    info.compress_type = ZIP_DEFLATED
    info.external_attr = 0o100644 << 16
    info.flag_bits |= 0x800
    return info


def _write_delivery_zip(
    output_path: Path,
    files: list[tuple[Path, str]],
    *,
    extra_payloads: dict[str, bytes] | None = None,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f"{output_path.stem}_",
        suffix=".zip",
        dir=output_path.parent,
        delete=False,
    ) as stream:
        temporary_path = Path(stream.name)
    try:
        with ZipFile(temporary_path, "w", compression=ZIP_DEFLATED) as package:
            for source, archive_name in sorted(files, key=lambda item: item[1]):
                normalized = _normalized_zip_name(archive_name)
                if (
                    PurePosixPath(normalized).name in DELIVERY_EXCLUDED_PARTS
                    or "__MACOSX" in PurePosixPath(normalized).parts
                    or PurePosixPath(normalized).name.startswith("._")
                ):
                    continue
                package.writestr(_zip_info(normalized), source.read_bytes())
            for name, payload in sorted((extra_payloads or {}).items()):
                package.writestr(_zip_info(name), payload)
        with ZipFile(temporary_path) as check:
            broken_member = check.testzip()
            if broken_member:
                raise ValueError(f"交付 ZIP 损坏: {broken_member}")
            names = check.namelist()
            if names != [_normalized_zip_name(name) for name in names]:
                raise ValueError(f"{output_path.name} 存在非 NFC 文件名")
            if any(
                "__MACOSX" in PurePosixPath(name).parts
                or PurePosixPath(name).name.startswith("._")
                or PurePosixPath(name).name == ".DS_Store"
                for name in names
            ):
                raise ValueError(f"{output_path.name} 包含 macOS 垃圾文件")
        temporary_path.replace(output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _code_snapshot_payloads(
    repo_root: Path,
    explicit_files: list[Path],
) -> dict[str, bytes]:
    completed = subprocess.run(
        ["git", "archive", "--format=tar", "HEAD"],
        cwd=repo_root,
        check=True,
        capture_output=True,
    )
    payloads: dict[str, bytes] = {}
    with tarfile.open(fileobj=BytesIO(completed.stdout), mode="r:") as archive:
        for member in archive.getmembers():
            if not member.isfile():
                continue
            stream = archive.extractfile(member)
            if stream is None:
                continue
            payloads[f"墨痕教育/{member.name}"] = stream.read()
    for path in explicit_files:
        if path.is_file():
            relative = path.resolve().relative_to(repo_root)
            payloads[f"墨痕教育/{relative.as_posix()}"] = path.read_bytes()
    return payloads


def package_deliverables(
    project_root: Path,
    repo_root: Path,
    question_results: list[DocumentResult],
    answer_results: list[AnswerResult],
    manifest_payload: bytes,
) -> dict[str, str]:
    delivery = project_root / "交付压缩包"
    ready_questions = [
        (Path(result.output), Path(result.output).name)
        for result in question_results
        if result.kind != "微点拓展专练"
    ]
    micro_questions = [
        (Path(result.output), Path(result.output).name)
        for result in question_results
        if result.kind == "微点拓展专练"
    ]
    cleaned_answers: list[tuple[Path, str]] = []
    for result in answer_results:
        answer_path = Path(result.output)
        cleaned_answers.append((answer_path, answer_path.name))
        status_path = answer_path.with_name(f"{answer_path.stem}_审核状态.json")
        cleaned_answers.append((status_path, status_path.name))

    paths = {
        "questions": delivery / "guanmei_geography_questions_windows.zip",
        "answers": delivery / "guanmei_geography_cleaned_answers_windows.zip",
        "micro": delivery / "guanmei_geography_micro_exercises_no_answers_windows.zip",
        "code": delivery / "mohen_education_guanmei_geography_code_windows.zip",
    }
    extras = {"guanmei_geography_delivery_manifest.json": manifest_payload}
    _write_delivery_zip(paths["questions"], ready_questions, extra_payloads=extras)
    _write_delivery_zip(paths["answers"], cleaned_answers, extra_payloads=extras)
    _write_delivery_zip(paths["micro"], micro_questions, extra_payloads=extras)

    code_payloads = _code_snapshot_payloads(
        repo_root,
        [
            repo_root / "tools/process_guanmei_geography.py",
            repo_root / "test_guanmei_geography_workflow.py",
        ],
    )
    code_payloads.update(extras)
    _write_delivery_zip(
        paths["code"],
        [],
        extra_payloads=code_payloads,
    )
    return {key: str(path) for key, path in paths.items()}


def process_project(project_root: str | Path, *, repo_root: str | Path) -> dict:
    project = Path(project_root).resolve()
    repo = Path(repo_root).resolve()
    source = project / "新高二课时分层作业.docx"
    if not source.is_file():
        raise FileNotFoundError(f"未找到总题目文档: {source}")
    source_hash_before = _sha256(source)

    question_results = split_question_document(source, project)
    answer_results = clean_answer_batch(project, question_results)
    offline_input_rehearsal = validate_offline_input_mapping(
        question_results,
        answer_results,
    )

    ready_questions = [
        result for result in question_results if result.kind != "微点拓展专练"
    ]
    if len(ready_questions) != 29 or len(answer_results) != 29:
        raise ValueError(
            f"交付数量异常: 题目 {len(ready_questions)}，答案 {len(answer_results)}"
        )
    question_total = sum(result.question_count for result in ready_questions)
    answer_total = sum(result.answer_count for result in answer_results)
    if question_total <= 0 or question_total != answer_total:
        raise ValueError(
            f"题目/答案总题数异常: 题目 {question_total}，答案 {answer_total}"
        )
    source_hash_after = _sha256(source)
    if source_hash_after != source_hash_before:
        raise ValueError("总题目源文档在处理期间发生变化")

    manifest = {
        "project": "莞美-高二-地理",
        "source_question": {
            "path": str(source),
            "sha256_before": source_hash_before,
            "sha256_after": source_hash_after,
            "unchanged": True,
        },
        "acceptance": {
            "ready_question_documents": len(ready_questions),
            "micro_documents_without_answers": sum(
                result.kind == "微点拓展专练" for result in question_results
            ),
            "cleaned_answer_documents": len(answer_results),
            "question_units": question_total,
            "answer_units": answer_total,
            "micro_question_units_without_answers": sum(
                result.question_count
                for result in question_results
                if result.kind == "微点拓展专练"
            ),
            "question_margins_twips": [PAGE_MARGIN_TWIPS] * 4,
            "question_margins_cm_display": 2.0,
            "question_body_policy": "章节范围内全部正文、表格、公式和图片原样保留",
            "windows_wps_status": "待用户在 Windows WPS 生产环境打开确认",
        },
        "questions": [asdict(result) for result in question_results],
        "answers": [asdict(result) for result in answer_results],
        "offline_input_rehearsal": offline_input_rehearsal,
    }
    manifest_payload = json.dumps(
        manifest,
        ensure_ascii=False,
        indent=2,
    ).encode("utf-8")
    delivery_dir = project / "交付压缩包"
    delivery_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = delivery_dir / "guanmei_geography_delivery_manifest.json"
    manifest_path.write_bytes(manifest_payload)
    archives = package_deliverables(
        project,
        repo,
        question_results,
        answer_results,
        manifest_payload,
    )
    manifest["archives"] = {
        key: {"path": value, "sha256": _sha256(Path(value))}
        for key, value in archives.items()
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="拆分莞美高二地理题目、清洗答案并生成 Windows 交付包"
    )
    parser.add_argument("project_root", type=Path, help="莞美-高二-地理资料目录")
    parser.add_argument(
        "--repo-root",
        type=Path,
        default=PROJECT_ROOT,
        help="墨痕教育代码仓库目录",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    manifest = process_project(args.project_root, repo_root=args.repo_root)
    acceptance = manifest["acceptance"]
    print(
        "处理完成："
        f"{acceptance['ready_question_documents']} 份可录入题目，"
        f"{acceptance['cleaned_answer_documents']} 份清洗答案，"
        f"{acceptance['question_units']} 道题；"
        "Windows WPS 仍待人工确认。"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
