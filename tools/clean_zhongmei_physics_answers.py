from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import asdict, dataclass, replace
from hashlib import sha256
import json
from pathlib import Path
import re
import sys
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from lxml import etree as ET

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from shared_core import (
    build_answer_units_from_docx,
    build_question_units_from_docx,
    build_review_report,
    export_review_report,
    initialize_review_status,
    map_answers,
    update_review_status,
)
from 墨痕快刀.config import SCIENCE_SCORE_CUE_PATTERN


DEFAULT_PROJECT_ROOT = Path(r"D:\墨痕教育题目\众美-高三-物理")
DEFAULT_REPORT_ROOT = Path(r"E:\CODEX.projection\墨痕教育-项目产物\众美-高三-物理")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
W = f"{{{W_NS}}}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
DOCUMENT_XML = "word/document.xml"

ET.register_namespace("w", W_NS)

ANSWER_MARKER_PATTERN = re.compile(
    r"^\s*(?:【答案】|答案)(?:(?:\s*[：:]\s*)|\s+|(?=$))"
)
ANALYSIS_MARKER_PATTERN = re.compile(
    r"^\s*(?:【解析】|解析)(?:(?:\s*[：:]\s*)|\s+|(?=$))"
)
INLINE_ANALYSIS_PATTERN = re.compile(
    r"(?:【解析】|解析)(?:(?:\s*[：:]\s*)|\s+|(?=$))"
)
QUESTION_SOURCE_START_PATTERN = re.compile(
    r"^\s*(?:[★☆/]\s*)?\d+\s*(?:[．.、]|[（(])"
)
CHOICE_ANSWER_PATTERN = re.compile(
    r"^[A-HＡ-Ｈ]+(?:\s*[,，、/]\s*[A-HＡ-Ｈ]+)*[。.]?$",
    re.IGNORECASE,
)
SCORE_CUE_PATTERN = re.compile(SCIENCE_SCORE_CUE_PATTERN)

RICH_CONTENT_TAGS = {
    f"{{{M_NS}}}oMath",
    f"{{{M_NS}}}oMathPara",
    f"{{{W_NS}}}drawing",
    f"{{{W_NS}}}pict",
    f"{{{W_NS}}}object",
    f"{{{A_NS}}}blip",
    f"{{{V_NS}}}imagedata",
}


@dataclass(frozen=True)
class NormalizationResult:
    source_path: Path
    output_path: Path
    answer_count: int
    inserted_analysis_placeholders: int
    rich_answers_moved_to_analysis: int
    inline_answer_analysis_count: int


@dataclass(frozen=True)
class PreflightPair:
    question_path: Path
    answer_path: Path
    output_path: Path
    relative_key: str
    question_count: int
    answer_count: int
    question_sha256: str
    answer_sha256: str


@dataclass(frozen=True)
class BatchResult:
    question_path: Path
    source_path: Path
    output_path: Path
    review_report_path: Path
    question_count: int
    answer_count: int
    inserted_analysis_placeholders: int
    rich_answers_moved_to_analysis: int
    inline_answer_analysis_count: int
    blocking_issue_count: int


@dataclass(frozen=True)
class _SourceBlock:
    answer_paragraph: object
    answer_children: tuple[object, ...]
    analysis_children: tuple[object, ...]
    answer_match_end: int
    inline_answer_text: str | None
    inline_analysis_text: str | None


def _paragraph_text(paragraph) -> str:
    return "".join(node.text or "" for node in paragraph.iter(f"{W}t"))


def _set_text(node, value: str) -> None:
    node.text = value
    if value.startswith(" ") or value.endswith(" "):
        node.set(XML_SPACE, "preserve")
    else:
        node.attrib.pop(XML_SPACE, None)


def _replace_text_prefix(paragraph, prefix_length: int, replacement: str) -> None:
    text_nodes = list(paragraph.iter(f"{W}t"))
    if not text_nodes:
        raise ValueError("答案或解析标签所在段落没有可编辑文本节点")

    cursor = 0
    replacement_written = False
    for node in text_nodes:
        value = node.text or ""
        start = cursor
        end = start + len(value)
        cursor = end

        if not replacement_written:
            suffix = value[max(0, prefix_length - start) :] if prefix_length < end else ""
            _set_text(node, replacement + suffix)
            replacement_written = True
            continue

        if end <= prefix_length:
            _set_text(node, "")
        elif start < prefix_length:
            _set_text(node, value[prefix_length - start :])


def _normalize_marker_run_format(paragraph, marker: str) -> None:
    marker_node = next(
        (
            node
            for node in paragraph.iter(f"{W}t")
            if (node.text or "").startswith(marker)
        ),
        None,
    )
    if marker_node is None:
        raise ValueError(f"标准化后的段落缺少标记: {marker}")

    run = marker_node.getparent()
    while run is not None and run.tag != f"{W}r":
        run = run.getparent()
    if run is None:
        raise ValueError(f"标记不在普通文本运行中，无法设置可见格式: {marker}")

    run_properties = run.find(f"{W}rPr")
    if run_properties is None:
        run_properties = ET.Element(f"{W}rPr")
        run.insert(0, run_properties)

    for hidden_tag in ("vanish", "webHidden", "specVanish"):
        for hidden_node in list(run_properties.findall(f"{W}{hidden_tag}")):
            run_properties.remove(hidden_node)

    run_fonts = run_properties.find(f"{W}rFonts")
    if run_fonts is None:
        run_fonts = ET.SubElement(run_properties, f"{W}rFonts")
    run_fonts.set(f"{W}ascii", "Times New Roman")
    run_fonts.set(f"{W}hAnsi", "Times New Roman")
    run_fonts.set(f"{W}eastAsia", "宋体")
    run_fonts.set(f"{W}cs", "Times New Roman")
    run_fonts.set(f"{W}hint", "eastAsia")

    color = run_properties.find(f"{W}color")
    if color is None:
        color = ET.SubElement(run_properties, f"{W}color")
    color.set(f"{W}val", "000000")
    color.attrib.pop(f"{W}themeColor", None)
    color.attrib.pop(f"{W}themeTint", None)
    color.attrib.pop(f"{W}themeShade", None)


def _new_text_paragraph(reference_paragraph, value: str):
    paragraph = ET.Element(f"{W}p")
    paragraph_properties = reference_paragraph.find(f"{W}pPr")
    if paragraph_properties is not None:
        paragraph.append(deepcopy(paragraph_properties))

    run = ET.SubElement(paragraph, f"{W}r")
    reference_run_properties = reference_paragraph.find(f".//{W}rPr")
    if reference_run_properties is not None:
        run.append(deepcopy(reference_run_properties))
    text = ET.SubElement(run, f"{W}t")
    _set_text(text, value)
    if value.startswith("解析："):
        _normalize_marker_run_format(paragraph, "解析：")
    return paragraph


def _has_rich_content(children: tuple[object, ...]) -> bool:
    return any(
        node.tag in RICH_CONTENT_TAGS
        for child in children
        for node in child.iter()
    )


def _looks_like_choice_answer(value: str) -> bool:
    return bool(CHOICE_ANSWER_PATTERN.fullmatch((value or "").strip()))


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _body_end_position(body, children: list[object]) -> int:
    section_properties = body.find(f"{W}sectPr")
    return children.index(section_properties) if section_properties is not None else len(children)


def _find_next_block_boundary(
    children: list[object],
    answer_position: int,
    next_answer_position: int,
    *,
    has_next_answer: bool,
) -> int:
    if not has_next_answer:
        return next_answer_position

    question_positions = []
    score_positions = []
    for position in range(answer_position + 1, next_answer_position):
        child = children[position]
        if child.tag != f"{W}p":
            continue
        text = _paragraph_text(child).strip()
        if QUESTION_SOURCE_START_PATTERN.match(text):
            question_positions.append(position)
        if SCORE_CUE_PATTERN.match(text):
            score_positions.append(position)

    boundary = question_positions[-1] if question_positions else next_answer_position
    preceding_score_positions = [position for position in score_positions if position <= boundary]
    if preceding_score_positions:
        boundary = min(boundary, preceding_score_positions[0])
    return boundary


def _extract_source_blocks(root, source_name: str) -> list[_SourceBlock]:
    body = root.find(f"{W}body")
    if body is None:
        raise ValueError(f"文档缺少正文节点: {source_name}")
    children = list(body)
    body_end = _body_end_position(body, children)

    answer_positions = []
    for position, child in enumerate(children[:body_end]):
        if child.tag != f"{W}p":
            continue
        text = _paragraph_text(child)
        match = ANSWER_MARKER_PATTERN.match(text)
        if match:
            answer_positions.append((position, child, match))
        elif re.match(r"^\s*(?:【答案|答案)", text):
            raise ValueError(f"{source_name} 存在未知答案标记: {text[:80]}")

    if not answer_positions:
        raise ValueError(f"{source_name} 未识别到任何答案标记")

    blocks = []
    for index, (answer_position, answer_paragraph, answer_match) in enumerate(answer_positions):
        has_next_answer = index + 1 < len(answer_positions)
        next_answer_position = (
            answer_positions[index + 1][0] if has_next_answer else body_end
        )
        boundary = _find_next_block_boundary(
            children,
            answer_position,
            next_answer_position,
            has_next_answer=has_next_answer,
        )

        source_text = _paragraph_text(answer_paragraph)
        payload = source_text[answer_match.end() :]
        inline_match = INLINE_ANALYSIS_PATTERN.search(payload)
        if inline_match and not _looks_like_choice_answer(payload[: inline_match.start()]):
            inline_match = None
        inline_answer_text = None
        inline_analysis_text = None
        if inline_match:
            inline_answer_text = payload[: inline_match.start()].strip()
            inline_analysis_text = payload[inline_match.end() :].strip()

        analysis_position = None
        for position in range(answer_position + 1, boundary):
            child = children[position]
            if child.tag != f"{W}p":
                continue
            if ANALYSIS_MARKER_PATTERN.match(_paragraph_text(child)):
                analysis_position = position
                break

        if inline_match and analysis_position is not None:
            raise ValueError(f"{source_name} 同一道答案同时存在内联解析和独立解析")

        answer_end = analysis_position if analysis_position is not None else boundary
        answer_children = tuple(children[answer_position:answer_end])
        analysis_children = (
            tuple(children[analysis_position:boundary])
            if analysis_position is not None
            else ()
        )
        blocks.append(
            _SourceBlock(
                answer_paragraph=answer_paragraph,
                answer_children=answer_children,
                analysis_children=analysis_children,
                answer_match_end=answer_match.end(),
                inline_answer_text=inline_answer_text,
                inline_analysis_text=inline_analysis_text,
            )
        )
    return blocks


def _normalized_analysis_children(block: _SourceBlock) -> list[object]:
    if block.inline_analysis_text is not None:
        return [
            _new_text_paragraph(
                block.answer_paragraph,
                f"解析：{block.inline_analysis_text}",
            )
        ]
    if not block.analysis_children:
        return []

    normalized = [deepcopy(child) for child in block.analysis_children]
    first_paragraph = normalized[0]
    if first_paragraph.tag != f"{W}p":
        raise ValueError("解析块首节点不是段落，无法标准化解析标记")
    match = ANALYSIS_MARKER_PATTERN.match(_paragraph_text(first_paragraph))
    if not match:
        raise ValueError("解析块首段缺少解析标记")
    _replace_text_prefix(first_paragraph, match.end(), "解析：")
    _normalize_marker_run_format(first_paragraph, "解析：")
    return normalized


def _render_block(block: _SourceBlock, question_id: int) -> tuple[list[object], bool, bool]:
    answer_payload = (
        block.inline_answer_text
        if block.inline_answer_text is not None
        else _paragraph_text(block.answer_paragraph)[block.answer_match_end :].strip()
    )
    rich_answer = _has_rich_content(block.answer_children)
    choice_answer = (
        not rich_answer
        and len(block.answer_children) == 1
        and _looks_like_choice_answer(answer_payload)
    )
    analysis_children = _normalized_analysis_children(block)
    inserted_placeholder = not analysis_children

    rendered: list[object] = []
    if choice_answer:
        if block.inline_answer_text is not None:
            rendered.append(
                _new_text_paragraph(block.answer_paragraph, f"{question_id}．{answer_payload}")
            )
        else:
            answer_paragraph = deepcopy(block.answer_paragraph)
            _replace_text_prefix(answer_paragraph, block.answer_match_end, f"{question_id}．")
            rendered.append(answer_paragraph)
        rendered.extend(
            analysis_children
            if analysis_children
            else [_new_text_paragraph(block.answer_paragraph, "解析： ")]
        )
        return rendered, inserted_placeholder, False

    rendered.append(_new_text_paragraph(block.answer_paragraph, f"{question_id}．"))
    if rich_answer:
        rendered.append(_new_text_paragraph(block.answer_paragraph, "答案： "))
    else:
        normalized_answer_children = [deepcopy(child) for child in block.answer_children]
        first_answer = normalized_answer_children[0]
        _replace_text_prefix(first_answer, block.answer_match_end, "答案：")
        rendered.extend(normalized_answer_children)

    rendered.extend(
        analysis_children
        if analysis_children
        else [_new_text_paragraph(block.answer_paragraph, "解析： ")]
    )
    if rich_answer:
        normalized_answer_children = [deepcopy(child) for child in block.answer_children]
        first_answer = normalized_answer_children[0]
        _replace_text_prefix(first_answer, block.answer_match_end, "答案：")
        rendered.extend(normalized_answer_children)
    return rendered, inserted_placeholder, rich_answer


def _write_docx_with_document_xml(
    source_path: Path,
    output_path: Path,
    document_xml: bytes,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if source_path.resolve() == output_path.resolve():
        raise ValueError("清洗输出不能覆盖源答案文档")

    with tempfile.NamedTemporaryFile(
        prefix=f"{output_path.stem}_",
        suffix=".docx",
        dir=output_path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)

    try:
        with ZipFile(source_path, "r") as source_zip, ZipFile(
            temporary_path,
            "w",
            compression=ZIP_DEFLATED,
        ) as target_zip:
            for info in source_zip.infolist():
                payload = document_xml if info.filename == DOCUMENT_XML else source_zip.read(info.filename)
                target_zip.writestr(info, payload)

        with ZipFile(temporary_path, "r") as check_zip:
            bad_member = check_zip.testzip()
            if bad_member:
                raise ValueError(f"清洗后的 docx 压缩包损坏: {bad_member}")
        Document(temporary_path)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def normalize_answer_docx(
    source_path: str | Path,
    output_path: str | Path,
    *,
    expected_question_count: int | None = None,
) -> NormalizationResult:
    source = Path(source_path)
    output = Path(output_path)
    if not source.is_file():
        raise FileNotFoundError(f"未找到答案文档: {source}")

    with ZipFile(source, "r") as package:
        root = ET.fromstring(package.read(DOCUMENT_XML))
    blocks = _extract_source_blocks(root, source.name)
    if expected_question_count is not None and len(blocks) != expected_question_count:
        raise ValueError(
            f"{source.name} 题答数量不一致: 题目 {expected_question_count}，答案 {len(blocks)}"
        )

    body = root.find(f"{W}body")
    section_properties = body.find(f"{W}sectPr") if body is not None else None
    if body is None:
        raise ValueError(f"文档缺少正文节点: {source.name}")
    for child in list(body):
        if child is not section_properties:
            body.remove(child)

    inserted_placeholders = 0
    rich_answers = 0
    inline_answers = 0
    for question_id, block in enumerate(blocks, 1):
        rendered, inserted_placeholder, rich_answer = _render_block(block, question_id)
        for child in rendered:
            if section_properties is None:
                body.append(child)
            else:
                body.insert(list(body).index(section_properties), child)
        inserted_placeholders += int(inserted_placeholder)
        rich_answers += int(rich_answer)
        inline_answers += int(block.inline_analysis_text is not None)

    document_xml = ET.tostring(
        root,
        encoding="utf-8",
        xml_declaration=True,
        standalone=True,
    )
    _write_docx_with_document_xml(source, output, document_xml)

    parsed_units = build_answer_units_from_docx(output, preserve_source_positions=True)
    if len(parsed_units) != len(blocks):
        output.unlink(missing_ok=True)
        raise ValueError(
            f"{source.name} 清洗后答案块数量异常: 期望 {len(blocks)}，实际 {len(parsed_units)}"
        )

    return NormalizationResult(
        source_path=source,
        output_path=output,
        answer_count=len(blocks),
        inserted_analysis_placeholders=inserted_placeholders,
        rich_answers_moved_to_analysis=rich_answers,
        inline_answer_analysis_count=inline_answers,
    )


def _pair_key(relative_path: Path, *, question: bool) -> str:
    filename = relative_path.name
    if question:
        filename = re.sub(r"^\s*【空白试卷】\s*", "", filename)
    return relative_path.with_name(filename).as_posix().casefold()


def _source_docx_files(root: Path, *, answer_files: bool) -> list[Path]:
    files = []
    for path in root.rglob("*.docx"):
        if path.name.startswith("~$") or path.stem.endswith("_已清洗"):
            continue
        if answer_files:
            files.append(path)
        elif "答案" not in path.relative_to(root).parts:
            files.append(path)
    return sorted(files, key=lambda path: str(path).casefold())


def preflight_project(project_root: str | Path) -> list[PreflightPair]:
    root = Path(project_root)
    answer_root = root / "答案"
    if not root.is_dir():
        raise FileNotFoundError(f"未找到众美物理项目目录: {root}")
    if not answer_root.is_dir():
        raise FileNotFoundError(f"未找到答案目录: {answer_root}")

    question_map = {
        _pair_key(path.relative_to(root), question=True): path
        for path in _source_docx_files(root, answer_files=False)
    }
    answer_map = {
        _pair_key(path.relative_to(answer_root), question=False): path
        for path in _source_docx_files(answer_root, answer_files=True)
    }
    if not question_map or not answer_map:
        raise ValueError("题目或答案文档为空，已停止整批预检")

    missing_answers = sorted(set(question_map) - set(answer_map))
    extra_answers = sorted(set(answer_map) - set(question_map))
    if missing_answers or extra_answers:
        detail = []
        if missing_answers:
            detail.append(f"缺答案 {len(missing_answers)} 份: {missing_answers[:3]}")
        if extra_answers:
            detail.append(f"多余答案 {len(extra_answers)} 份: {extra_answers[:3]}")
        raise ValueError("题答文件配对失败: " + "；".join(detail))

    pairs = []
    for key in sorted(question_map):
        question_path = question_map[key]
        answer_path = answer_map[key]
        question_units = build_question_units_from_docx(question_path)
        with ZipFile(answer_path, "r") as package:
            answer_root_xml = ET.fromstring(package.read(DOCUMENT_XML))
        answer_count = len(_extract_source_blocks(answer_root_xml, answer_path.name))
        question_count = len(question_units)
        if question_count != answer_count:
            raise ValueError(
                f"{key} 题答数量不一致: 题目 {question_count}，答案 {answer_count}"
            )
        pairs.append(
            PreflightPair(
                question_path=question_path,
                answer_path=answer_path,
                output_path=answer_path.with_name(f"{answer_path.stem}_已清洗.docx"),
                relative_key=key,
                question_count=question_count,
                answer_count=answer_count,
                question_sha256=_sha256(question_path),
                answer_sha256=_sha256(answer_path),
            )
        )
    return pairs


def _questions_for_review(question_path: Path):
    return [
        replace(unit, question_id=str(index), subquestions=[])
        if unit.question_type == "choice" and unit.subquestions
        else replace(unit, question_id=str(index))
        for index, unit in enumerate(build_question_units_from_docx(question_path), 1)
    ]


def _write_preflight_report(pairs: list[PreflightPair], report_root: Path) -> Path:
    report_root.mkdir(parents=True, exist_ok=True)
    report_path = report_root / "众美高三物理_整批预检.json"
    payload = {
        "project": "众美-高三-物理",
        "document_count": len(pairs),
        "question_count": sum(pair.question_count for pair in pairs),
        "source_immutable": True,
        "pairs": [
            {
                **asdict(pair),
                "question_path": str(pair.question_path.resolve()),
                "answer_path": str(pair.answer_path.resolve()),
                "output_path": str(pair.output_path.resolve()),
            }
            for pair in pairs
        ],
    }
    report_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return report_path


def clean_project(
    project_root: str | Path,
    *,
    report_root: str | Path = DEFAULT_REPORT_ROOT,
    overwrite: bool = False,
) -> list[BatchResult]:
    pairs = preflight_project(project_root)
    report_path_root = Path(report_root)
    _write_preflight_report(pairs, report_path_root)

    existing_outputs = [pair.output_path for pair in pairs if pair.output_path.exists()]
    if existing_outputs and not overwrite:
        raise FileExistsError(
            f"已有 {len(existing_outputs)} 份清洗结果；如需重建请显式使用 --overwrite"
        )

    answer_root = Path(project_root) / "答案"
    results = []
    for pair in pairs:
        normalization = normalize_answer_docx(
            pair.answer_path,
            pair.output_path,
            expected_question_count=pair.question_count,
        )
        answer_units = build_answer_units_from_docx(
            pair.output_path,
            preserve_source_positions=True,
        )
        question_units = _questions_for_review(pair.question_path)
        mapped_units = map_answers(question_units, answer_units)
        report = build_review_report(pair.output_path.name, question_units, mapped_units)

        relative_parent = pair.answer_path.relative_to(answer_root).parent
        review_dir = report_path_root / "审核清单" / relative_parent
        review_dir.mkdir(parents=True, exist_ok=True)
        review_report_path = review_dir / f"{pair.output_path.stem}_审核清单.md"
        export_review_report(report, review_report_path)
        initialize_review_status(
            pair.output_path,
            report_path=str(review_report_path),
            report=report,
        )
        blocking_issue_count = sum(
            1 for issue in report.issues if issue.severity == "error"
        )
        update_review_status(
            pair.output_path,
            status="approved" if blocking_issue_count == 0 else "rejected",
            reviewer="system",
            note=(
                "众美高三物理答案自动检查通过"
                if blocking_issue_count == 0
                else f"自动检查发现 {blocking_issue_count} 个高风险问题"
            ),
        )
        results.append(
            BatchResult(
                question_path=pair.question_path,
                source_path=pair.answer_path,
                output_path=pair.output_path,
                review_report_path=review_report_path,
                question_count=pair.question_count,
                answer_count=normalization.answer_count,
                inserted_analysis_placeholders=normalization.inserted_analysis_placeholders,
                rich_answers_moved_to_analysis=normalization.rich_answers_moved_to_analysis,
                inline_answer_analysis_count=normalization.inline_answer_analysis_count,
                blocking_issue_count=blocking_issue_count,
            )
        )
    return results


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="众美高三物理整批只读预检、答案富文本清洗与审核门禁生成",
    )
    parser.add_argument(
        "--project-root",
        default=str(DEFAULT_PROJECT_ROOT),
        help="众美高三物理项目根目录",
    )
    parser.add_argument(
        "--report-root",
        default=str(DEFAULT_REPORT_ROOT),
        help="预检和审核报告输出目录（默认 E 盘）",
    )
    parser.add_argument("--preflight-only", action="store_true", help="只做整批预检，不写清洗文档")
    parser.add_argument("--overwrite", action="store_true", help="显式覆盖既有 _已清洗.docx")
    return parser


def main() -> int:
    args = _build_parser().parse_args()
    pairs = preflight_project(args.project_root)
    preflight_report = _write_preflight_report(pairs, Path(args.report_root))
    print(f"预检文件: {len(pairs)}")
    print(f"预检题数: {sum(pair.question_count for pair in pairs)}")
    print(f"预检报告: {preflight_report.resolve()}")
    if args.preflight_only:
        return 0

    results = clean_project(
        args.project_root,
        report_root=args.report_root,
        overwrite=args.overwrite,
    )
    approved = sum(result.blocking_issue_count == 0 for result in results)
    print(f"清洗文件: {len(results)}")
    print(f"答案总数: {sum(result.answer_count for result in results)}")
    print(f"补空解析占位: {sum(result.inserted_analysis_placeholders for result in results)}")
    print(f"富文本答案转入解析: {sum(result.rich_answers_moved_to_analysis for result in results)}")
    print(f"内联答案解析拆分: {sum(result.inline_answer_analysis_count for result in results)}")
    print(f"自动检查通过: {approved}/{len(results)}")
    return 0 if approved == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
